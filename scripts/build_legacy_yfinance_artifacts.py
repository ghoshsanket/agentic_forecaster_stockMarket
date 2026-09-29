#!/usr/bin/env python3
"""Build and validate the LEGACY Yahoo Finance dataset artifacts.

Consumes the canonical Parquet files written by ``download_yfinance_daily.py``
for the USER_SUPPLIED_LEGACY_NIFTY50_UNIVERSE and produces, under
``$AGENTIC_YFINANCE_LEGACY_ROOT``:

    workbooks/adjusted/Nifty-50-Legacy-Historical-Securities.xlsx
    workbooks/unadjusted/Nifty-50-Legacy-Historical-Securities.xlsx
    workbooks/adjusted/Nifty-50.xlsx        (original-convention copy)
    workbooks/unadjusted/Nifty-50.xlsx      (ORIGINAL-RECOLLECTION CANDIDATE)

    metadata/legacy_coverage.csv
    metadata/security_lineage.csv
    metadata/successor_lineage_plan.csv
    metadata/modern_vs_legacy_universe.csv
    metadata/sheet_name_map.csv
    metadata/adjusted_vs_unadjusted.csv
    metadata/data_quality.json
    metadata/coverage_summary.json

    manifests/legacy_yfinance_manifest.json
    manifests/manifest_sha256.csv

    corporate_history/<legacy_label>.md

Rules enforced here:

* ONE sheet per user-supplied label, named with the EXACT label.
* A label whose security is unavailable gets an EMPTY sheet holding only the
  canonical header.  No rows are invented.
* No commentary rows are ever placed inside an OHLCV table.
* Successor prices are NEVER appended.
* Every workbook is read back with openpyxl and compared against the canonical
  data before it is accepted.

Usage:
    python scripts/build_legacy_yfinance_artifacts.py
"""

from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import platform
import sys
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
import pandas as pd

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "src"))

_spec = importlib.util.spec_from_file_location(
    "download_yfinance_daily", REPO_ROOT / "scripts" / "download_yfinance_daily.py")
dl = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(dl)

CANONICAL = dl.CANONICAL_COLUMNS
VARIANTS = dl.VARIANTS
UNIVERSE_CFG = "configs/nifty50_legacy_user_supplied.yaml"
LINEAGE_CFG = "configs/legacy_security_lineage.yaml"

LINEAGE_COLUMNS = [
    "legacy_label", "historical_company_name", "current_or_final_company_name",
    "event_type", "primary_yahoo_candidate", "additional_yahoo_candidates",
    "same_legal_security", "successor_security", "event_date",
    "primary_download_policy", "confidence", "evidence_notes",
]
COVERAGE_COLUMNS = [
    "legacy_label", "historical_company_name", "accepted_yahoo_symbol",
    "identity_type", "event_type", "adjusted_status", "unadjusted_status",
    "first_date", "last_date", "rows", "ceased_to_exist_date", "successor",
    "date_anomaly", "notes",
]
SUCCESSOR_COLUMNS = [
    "legacy_label", "successor", "event", "event_date",
    "exchange_ratio_if_known", "continuation_possible", "evidence", "notes",
]


def legacy_root() -> Path:
    import os
    root = os.environ.get("AGENTIC_YFINANCE_LEGACY_ROOT")
    if root:
        return Path(root)
    return Path(os.environ["AGENTIC_DATA_ROOT"]) / "yfinance_legacy_nifty50_2000_2025"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def load_context() -> tuple[list[str], dict, dict]:
    labels = dl.load_universe(REPO_ROOT, UNIVERSE_CFG)
    lineage = dl.load_lineage(REPO_ROOT, LINEAGE_CFG)
    results_path = legacy_root() / "metadata" / "download_results.json"
    results = json.loads(results_path.read_text()) if results_path.is_file() else {}
    return labels, lineage, results


def read_series(root: Path, label: str, variant: str) -> pd.DataFrame | None:
    p = root / variant / "parquet" / f"{dl.safe_filename(label)}.parquet"
    return pd.read_parquet(p) if p.is_file() else None


# ------------------------------------------------------------------ workbook

def build_workbook(root: Path, variant: str, labels: list[str], filename: str) -> dict:
    from openpyxl import Workbook, load_workbook

    path = root / "workbooks" / variant / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    expected: dict[str, pd.DataFrame | None] = {}
    for label in labels:
        sheet = dl.legacy_sheet_name(label)
        expected[sheet] = read_series(root, label, variant)
        ws = wb.create_sheet(title=sheet)
        # Header is ALWAYS written, even for an unavailable security.
        ws.append(CANONICAL)
        df = expected[sheet]
        if df is not None and not df.empty:
            for rec in df.itertuples(index=False):
                ws.append([rec.Date.to_pydatetime(), float(rec.Open), float(rec.High),
                           float(rec.Low), float(rec.Close), int(rec.Volume)])
    wb.save(path)

    # ---- round-trip verification ----
    problems: list[str] = []
    check = load_workbook(path, read_only=True, data_only=True)
    if sorted(check.sheetnames) != sorted(expected):
        problems.append(f"sheet set mismatch: {sorted(check.sheetnames)}")
    empty_sheets: list[str] = []
    for label in labels:
        sheet = dl.legacy_sheet_name(label)
        if sheet not in check.sheetnames:
            problems.append(f"{label}: sheet missing")
            continue
        rows = check[sheet].iter_rows(values_only=True)
        header = next(rows, None)
        if list(header or []) != CANONICAL:
            problems.append(f"{sheet}: header {header} != {CANONICAL}")
            continue
        body = list(rows)
        df = expected[sheet]
        if df is None or df.empty:
            if body:
                problems.append(f"{sheet}: expected EMPTY but has {len(body)} rows")
            else:
                empty_sheets.append(label)
            continue
        if len(body) != len(df):
            problems.append(f"{sheet}: {len(body)} rows != {len(df)}")
            continue
        for i, (rec, got) in enumerate(zip(df.itertuples(index=False), body)):
            exp = (rec.Date.to_pydatetime(), float(rec.Open), float(rec.High),
                   float(rec.Low), float(rec.Close), int(rec.Volume))
            if got[0] != exp[0]:
                problems.append(f"{sheet}[{i}]: date {got[0]} != {exp[0]}")
                break
            for j in range(1, 6):
                if abs(float(got[j]) - exp[j]) > 1e-9:
                    problems.append(f"{sheet}[{i}]: {CANONICAL[j]} {got[j]} != {exp[j]}")
                    break
            else:
                continue
            break
    check.close()
    return {
        "path": str(path.relative_to(root)), "sheets": len(expected),
        "bytes": path.stat().st_size, "round_trip_ok": not problems,
        "round_trip_problems": problems[:20],
        "empty_sheets": empty_sheets,
    }


# ------------------------------------------------------------------- reports

def build_coverage(root: Path, labels: list[str], lineage: dict,
                   results: dict) -> tuple[pd.DataFrame, dict]:
    anomaly = {"POND'S"}
    rows = []
    for label in labels:
        e = lineage[label]
        pol = e.get("primary_download_policy")
        res = results.get(label, {})
        adj = res.get("adjusted", {})
        unadj = res.get("unadjusted", {})
        df = read_series(root, label, "adjusted")
        successor = e.get("successor_security") or ""
        notes = []
        if pol == "do_not_download":
            notes.append("no same-security Yahoo history; successor NOT appended")
        if e.get("same_legal_security") is False and pol != "do_not_download":
            notes.append("series contains a corporate-action discontinuity")
        if e.get("event_type") == "demerger" and pol == "download_primary_flagged":
            notes.append(f"STRUCTURAL BREAK at {e.get('event_date')}")
        if e.get("confidence") in ("medium", "low"):
            notes.append(f"confidence={e.get('confidence')}")
        rows.append({
            "legacy_label": label,
            "historical_company_name": e.get("historical_company_name"),
            "accepted_yahoo_symbol": e.get("primary_yahoo_candidate") or "",
            "identity_type": "same_legal_security" if e.get("same_legal_security")
                             else "different_legal_entity",
            "event_type": e.get("event_type"),
            "adjusted_status": adj.get("status", "unavailable_from_yahoo"),
            "unadjusted_status": unadj.get("status", "unavailable_from_yahoo"),
            "first_date": (str(df["Date"].min().date()) if df is not None and not df.empty else ""),
            "last_date": (str(df["Date"].max().date()) if df is not None and not df.empty else ""),
            "rows": len(df) if df is not None else 0,
            "ceased_to_exist_date": (e.get("event_date")
                                     if pol == "do_not_download" else ""),
            "successor": successor,
            "date_anomaly": "DATE_ANOMALY_TRUE" if label in anomaly else "FALSE",
            "notes": "; ".join(notes),
        })
    cov = pd.DataFrame(rows)[COVERAGE_COLUMNS]
    ok = cov[(cov.adjusted_status == "complete") & (cov.unadjusted_status == "complete")]
    summary = {
        "requested": len(labels),
        "downloaded_both_variants": len(ok),
        "unavailable_from_yahoo": int((cov.adjusted_status != "complete").sum()),
        "total_rows_adjusted": int(cov.rows.sum()),
        "earliest_date": cov[cov.first_date != ""].first_date.min() if len(ok) else "",
        "latest_date": cov[cov.last_date != ""].last_date.max() if len(ok) else "",
        "by_event_type": {k: int(v) for k, v in cov.event_type.value_counts().items()},
        "date_anomaly_labels": sorted(anomaly),
        "securities_starting_after_2000_01_03": sorted(
            cov[(cov.first_date != "") & (cov.first_date > "2000-01-03")].legacy_label.tolist()),
    }
    return cov, summary


def build_successor_plan(labels: list[str], lineage: dict) -> pd.DataFrame:
    ratios = {
        "BURROUGHS": "14 GSK : 10 BWIL", "COCHINREFN": "4 BPCL : 9 KRL",
        "IBP": "110 IOC : 100 IBP", "ICICI": "1 ICICI Bank : 2 ICICI",
        "HDFC": "42 HDFC Bank : 25 HDFC", "POND'S": "3 HLL : 4 PIL",
        "RANBAXY": "0.8 SUNPHARMA : 1 RANBAXY", "RHONE-POUL": "7 NPIL : 4 RPIL",
        "SATYAMCOMP": "2 TECHM : 17 Mahindra Satyam",
    }
    rows = []
    for label in labels:
        e = lineage[label]
        succ = e.get("successor_security")
        if not succ:
            continue
        ev = e.get("event_type")
        rows.append({
            "legacy_label": label,
            "successor": succ,
            "event": ev,
            "event_date": e.get("event_date"),
            "exchange_ratio_if_known": ratios.get(label, "not established"),
            "continuation_possible": "NO - not used in this dataset",
            "evidence": (e.get("evidence_notes") or "").strip()[:300],
            "notes": "Recorded for a FUTURE OPTIONAL experiment only. Successor "
                     "prices are NOT in the canonical workbook and must not be "
                     "stitched onto the historical series.",
        })
    return pd.DataFrame(rows)[SUCCESSOR_COLUMNS]


def build_universe_comparison() -> pd.DataFrame:
    """Compare the modern and legacy universes WITHOUT overwriting either.

    ``renamed_equivalent`` means the two labels differ but the modern label
    resolves to the same Yahoo symbol as the legacy label's verified
    same-security symbol, i.e. they denote the same listed issuer.
    """
    modern = dl.load_universe(REPO_ROOT, None)
    legacy = dl.load_universe(REPO_ROOT, UNIVERSE_CFG)
    lin = dl.load_lineage(REPO_ROOT, LINEAGE_CFG)

    modern_sym = {m: dl.yahoo_symbol(m) for m in modern}
    legacy_sym = {l: dl.yahoo_symbol(l, lin) for l in legacy}

    rows: list[dict] = []
    matched_modern: set[str] = set()
    matched_legacy: set[str] = set()

    # 1. exact label overlap
    for label in legacy:
        if label in modern:
            rows.append({"legacy_label": label, "modern_label": label,
                         "relationship": "exact",
                         "legacy_yahoo_symbol": legacy_sym[label],
                         "modern_yahoo_symbol": modern_sym[label]})
            matched_legacy.add(label)
            matched_modern.add(label)

    # 2. renamed-equivalent: different labels, same listed issuer
    for label in legacy:
        if label in matched_legacy:
            continue
        lsym = legacy_sym[label]
        if not lsym:
            continue
        for modern_label in modern:
            if modern_label in matched_modern:
                continue
            if modern_sym[modern_label] == lsym:
                rows.append({"legacy_label": label, "modern_label": modern_label,
                             "relationship": "renamed_equivalent",
                             "legacy_yahoo_symbol": lsym,
                             "modern_yahoo_symbol": modern_sym[modern_label]})
                matched_legacy.add(label)
                matched_modern.add(modern_label)
                break

    # 3. legacy-only
    for label in legacy:
        if label not in matched_legacy:
            rows.append({"legacy_label": label, "modern_label": "",
                         "relationship": "legacy_only",
                         "legacy_yahoo_symbol": legacy_sym[label],
                         "modern_yahoo_symbol": ""})

    # 4. modern-only
    for label in modern:
        if label not in matched_modern:
            rows.append({"legacy_label": "", "modern_label": label,
                         "relationship": "modern_only",
                         "legacy_yahoo_symbol": "",
                         "modern_yahoo_symbol": modern_sym[label]})

    return pd.DataFrame(rows)


def build_adj_vs_unadj(root: Path, labels: list[str]) -> pd.DataFrame:
    rows = []
    for label in labels:
        a = read_series(root, label, "adjusted")
        u = read_series(root, label, "unadjusted")
        if a is None or u is None:
            rows.append({"legacy_label": label, "adjusted_rows": 0, "unadjusted_rows": 0,
                         "adjusted_first": "", "adjusted_last": "", "unadjusted_first": "",
                         "unadjusted_last": "", "row_counts_equal": "", "split_like_jumps": "",
                         "return_correlation": "", "max_abs_close_diff": ""})
            continue
        m = a.merge(u, on="Date", suffixes=("_a", "_u"))
        ra, ru = m.Close_a.pct_change(), m.Close_u.pct_change()
        corr = float(ra.corr(ru)) if len(m) > 3 else np.nan
        rows.append({
            "legacy_label": label,
            "adjusted_rows": len(a), "unadjusted_rows": len(u),
            "adjusted_first": str(a.Date.min().date()), "adjusted_last": str(a.Date.max().date()),
            "unadjusted_first": str(u.Date.min().date()), "unadjusted_last": str(u.Date.max().date()),
            "row_counts_equal": bool(len(a) == len(u)),
            "split_like_jumps": int((m.Close_u.pct_change().abs() > 0.25).sum()),
            "return_correlation": None if pd.isna(corr) else round(corr, 6),
            "max_abs_close_diff": round(float((m.Close_a - m.Close_u).abs().max()), 4),
        })
    return pd.DataFrame(rows)


def build_quality(root: Path, labels: list[str]) -> dict:
    out: dict = {"definition": {
        "checks": ["ascending Date", "unique Date", "non-null OHLCV", "non-negative volume",
                   "valid OHLC relations", "within requested bounds"],
        "large_jumps": "recorded in coverage metadata, never deleted",
        "note": "vendor values reported as-is; nothing repaired or forward-filled"}}
    for variant in VARIANTS:
        issues, jumps = [], 0
        for label in labels:
            df = read_series(root, label, variant)
            if df is None or df.empty:
                continue
            p = []
            if not df.Date.is_monotonic_increasing:
                p.append("not_ascending")
            if df.Date.duplicated().any():
                p.append("duplicate_dates")
            if df[CANONICAL].isna().any().any():
                p.append("nulls")
            if (df.Volume < 0).any():
                p.append("negative_volume")
            v = df.dropna(subset=["Open", "High", "Low", "Close"])
            if (v.High < v.Low).any() or (v.High < v.Close).any() or (v.Low > v.Close).any():
                p.append("ohlc_relation")
            if df.Date.min() < pd.Timestamp("2000-01-01"):
                p.append("before_start")
            if df.Date.max() >= pd.Timestamp("2026-01-01"):
                p.append("on_or_after_end")
            if p:
                issues.append({"label": label, "issues": p})
            jumps += int((df.Close.pct_change().abs() > dl.JUMP_THRESHOLD).sum())
        out[variant] = {"securities_with_issues": issues, "large_jumps": jumps}
    return out


def write_corporate_history(root: Path, labels: list[str], lineage: dict) -> None:
    d = root / "corporate_history"
    d.mkdir(parents=True, exist_ok=True)
    for label in labels:
        e = lineage[label]
        pol = e.get("primary_download_policy")
        acquired = pol == "do_not_download"
        body = [
            f"# {label} — corporate identity record", "",
            f"- **Historical name:** {e.get('historical_company_name')}",
            f"- **Current / final name:** {e.get('current_or_final_company_name')}",
            f"- **Event type:** `{e.get('event_type')}`",
            f"- **Event date:** {e.get('event_date') or 'n/a'}",
            f"- **Same legal security:** {e.get('same_legal_security')}",
            f"- **Accepted Yahoo symbol:** `{e.get('primary_yahoo_candidate') or 'NONE — do_not_download'}`",
            f"- **Successor (recorded only, never appended):** {e.get('successor_security') or 'n/a'}",
            f"- **Confidence:** {e.get('confidence')}",
            f"- **Download policy:** `{pol}`", "",
            "## Evidence", "", (e.get("evidence_notes") or "").strip(), "",
        ]
        if acquired:
            body += [
                "## Policy applied", "",
                "This security no longer exists as a legal entity. **No successor price",
                "is appended, no forward-fill is performed and no synthetic pre-listing",
                "observations are created.** The canonical sheet for this label is",
                "intentionally empty and holds only the canonical header. The successor",
                "is recorded in `metadata/successor_lineage_plan.csv` for a future",
                "optional experiment that has NOT been performed.", "",
            ]
        (d / f"{dl.safe_filename(label)}.md").write_text("\n".join(body), encoding="utf-8")


def build_manifest(root: Path, labels: list[str], lineage_cfg: Path) -> dict:
    entries = []
    for variant in VARIANTS:
        for sub in ("csv", "parquet", "source_snapshots"):
            for f in sorted((root / variant / sub).iterdir()):
                if f.is_file():
                    entries.append({"path": str(f.relative_to(root)),
                                    "bytes": f.stat().st_size, "sha256": sha256(f)})
        for f in sorted((root / "workbooks" / variant).glob("*.xlsx")):
            entries.append({"path": str(f.relative_to(root)),
                            "bytes": f.stat().st_size, "sha256": sha256(f)})
    pd.DataFrame(entries).to_csv(root / "manifests" / "manifest_sha256.csv", index=False)

    uni_hash = sha256(REPO_ROOT / UNIVERSE_CFG)
    lin_hash = sha256(lineage_cfg)
    try:
        import yfinance
        yf_ver = yfinance.__version__
    except Exception as exc:
        yf_ver = f"unavailable ({type(exc).__name__})"
    return {
        "generated_at_utc": datetime.now(UTC).isoformat(),
        "universe_id": "USER_SUPPLIED_LEGACY_NIFTY50_UNIVERSE",
        "yfinance_version": yf_ver,
        "python": platform.python_version(),
        "platform": platform.platform(),
        "parameters": {"start": "2000-01-01", "end": "2026-01-01", "end_exclusive": True,
                       "interval": "1d", "actions": False, "repair": False, "prepost": False,
                       "threads": False, "adjusted": {"auto_adjust": True},
                       "unadjusted": {"auto_adjust": False}},
        "user_supplied_universe_sha256": uni_hash,
        "lineage_registry_sha256": lin_hash,
        "successor_stitching": False,
        "files": entries,
        "file_count": len(entries),
        "total_bytes": int(sum(e["bytes"] for e in entries)),
    }


def main() -> int:
    root = legacy_root()
    dl.ensure_layout(root)
    labels, lineage, results = load_context()
    print(f"legacy root : {root}")
    print(f"labels      : {len(labels)}")

    # ---- metadata: lineage + sheet map ----
    mdir = root / "metadata"
    with (mdir / "security_lineage.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=LINEAGE_COLUMNS)
        w.writeheader()
        for label in labels:
            e = lineage[label]
            row = {k: e.get(k, "") for k in LINEAGE_COLUMNS}
            row["additional_yahoo_candidates"] = ";".join(
                e.get("additional_yahoo_candidates") or [])
            w.writerow(row)
    pd.DataFrame([{
        "legacy_label": l,
        "canonical_filename_stem": dl.safe_filename(l),
        "csv_filename": f"{dl.safe_filename(l)}.csv",
        "parquet_filename": f"{dl.safe_filename(l)}.parquet",
        "excel_sheet_name": dl.legacy_sheet_name(l),
        "sheet_name_equals_legacy_label": dl.legacy_sheet_name(l) == l,
    } for l in labels]).to_csv(mdir / "sheet_name_map.csv", index=False)

    cov, cov_summary = build_coverage(root, labels, lineage, results)
    cov.to_csv(mdir / "legacy_coverage.csv", index=False)
    (mdir / "coverage_summary.json").write_text(json.dumps(cov_summary, indent=2))
    build_successor_plan(labels, lineage).to_csv(
        mdir / "successor_lineage_plan.csv", index=False)
    build_universe_comparison().to_csv(mdir / "modern_vs_legacy_universe.csv", index=False)
    build_adj_vs_unadj(root, labels).to_csv(mdir / "adjusted_vs_unadjusted.csv", index=False)
    (mdir / "data_quality.json").write_text(
        json.dumps(build_quality(root, labels), indent=2))
    write_corporate_history(root, labels, lineage)

    # ---- workbooks ----
    books = {}
    for variant in VARIANTS:
        books[f"{variant}_primary"] = build_workbook(
            root, variant, labels, "Nifty-50-Legacy-Historical-Securities.xlsx")
        books[f"{variant}_original_convention"] = build_workbook(
            root, variant, labels, "Nifty-50.xlsx")
    for name, info in books.items():
        print(f"workbook {name:34s} sheets={info['sheets']:3d} "
              f"round_trip_ok={info['round_trip_ok']} bytes={info['bytes']:>10,}")
        if info["round_trip_problems"]:
            print("   PROBLEMS:", info["round_trip_problems"][:3])

    manifest = build_manifest(root, labels, REPO_ROOT / LINEAGE_CFG)
    (root / "manifests" / "legacy_yfinance_manifest.json").write_text(
        json.dumps(manifest, indent=2))

    print(f"\ncoverage : {cov_summary['downloaded_both_variants']}/{cov_summary['requested']} "
          f"downloaded in both variants, {cov_summary['unavailable_from_yahoo']} unavailable")
    print(f"rows     : {cov_summary['total_rows_adjusted']:,} adjusted (unadjusted equal)")
    print(f"events   : {cov_summary['by_event_type']}")
    print(f"manifest : {manifest['file_count']} files, {manifest['total_bytes']:,} bytes")
    print(f"universe sha256: {manifest['user_supplied_universe_sha256'][:16]}...")
    print(f"lineage  sha256: {manifest['lineage_registry_sha256'][:16]}...")
    ok = all(b["round_trip_ok"] for b in books.values())
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
