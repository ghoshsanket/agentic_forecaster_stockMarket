"""Tests for the paper-snapshot universe and the pre-performance cleanup.

Covers:

* the canonical ``AGENTIC_OUTPUT_ROOT`` default (``output/``, not ``outputs/``)
* the paper-snapshot universe config, its provenance metadata, and the
  2025 constituent changes
* the DUMMYTATAM exclusion and the TMPV / Tata Motors treatment
* the three-universe comparison covering the paper's example tickers
* the paper-snapshot performance configs (100-epoch budget, not run)

Tests that need the downloaded paper-snapshot dataset are opt-in; the rest run
in a clean checkout.
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest
import yaml

REPO_ROOT = Path(__file__).resolve().parents[2]
PAPER_CFG = "configs/nifty50_paper_snapshot_2025_11_04.yaml"
MODERN_CFG = "configs/nifty50.yaml"
LEGACY_CFG = "configs/nifty50_legacy_user_supplied.yaml"

PAPER_EXAMPLES = ["RELIANCE", "TCS", "INFY", "HDFCBANK", "ITC"]
# Official NSE Indices 2025 changes, applied to reach the 2025-11-04 snapshot.
EXITS_2025 = ["BPCL", "BRITANNIA", "HEROMOTOCO", "INDUSINDBK"]
ENTRANTS_2025 = ["JIOFIN", "ZOMATO", "INDIGO", "MAXHEALTH"]


def _paper_root() -> Path:
    root = os.environ.get("AGENTIC_YFINANCE_PAPER_SNAPSHOT_ROOT")
    if root:
        return Path(root)
    from agentic_forecaster.config import _env_defaults
    data = Path(_env_defaults()["AGENTIC_DATA_ROOT"])
    return data / "yfinance_paper_snapshot_2025_11_04"


PAPER_ROOT = _paper_root()
HAS_PAPER_DATASET = (PAPER_ROOT / "unadjusted" / "csv").is_dir()
requires_paper_dataset = pytest.mark.skipif(
    not HAS_PAPER_DATASET, reason="paper-snapshot dataset is not present in this workspace")


def _load(path: str) -> dict:
    return yaml.safe_load((REPO_ROOT / path).read_text())


# ------------------------------------------------- canonical output root

def test_output_root_default_is_output_singular(monkeypatch):
    """The canonical workspace root is $RESEARCH_ROOT/output/agentic-forecaster."""
    from agentic_forecaster.config import _env_defaults
    root = str(REPO_ROOT.parents[1])
    for var in list(os.environ):
        if var.startswith("AGENTIC_"):
            monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv("RESEARCH_ROOT", root)
    assert _env_defaults()["AGENTIC_OUTPUT_ROOT"] == f"{root}/output/agentic-forecaster"


def test_output_root_env_still_overrides(monkeypatch):
    from agentic_forecaster.config import get_env_roots
    monkeypatch.setenv("AGENTIC_OUTPUT_ROOT", "/tmp/explicit-output")
    assert get_env_roots()["AGENTIC_OUTPUT_ROOT"] == "/tmp/explicit-output"


def test_docs_reference_the_canonical_output_root():
    for doc in ("docs/REPRODUCIBILITY.md", "docs/STORAGE_LAYOUT.md"):
        text = (REPO_ROOT / doc).read_text()
        assert "output/agentic-forecaster" in text, doc
        assert "outputs/agentic-forecaster" not in text, doc


def test_output_root_stays_under_research_root(monkeypatch):
    from agentic_forecaster.config import _env_defaults
    root = str(REPO_ROOT.parents[1])
    for var in list(os.environ):
        if var.startswith("AGENTIC_"):
            monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv("RESEARCH_ROOT", root)
    assert _env_defaults()["AGENTIC_OUTPUT_ROOT"].startswith(root)


# ------------------------------------------- paper-snapshot universe config

def test_paper_snapshot_config_metadata():
    cfg = _load(PAPER_CFG)
    assert cfg["universe_id"] == "PAPER_FIXED_NIFTY50_2025_11_04_CANDIDATE"
    assert cfg["status"] == "RECONSTRUCTED_FROM_OFFICIAL_NSE_EVIDENCE"
    assert str(cfg["snapshot_date"]) == "2025-11-04"
    assert cfg["requested_download_start"] == "2000-01-01"
    assert cfg["requested_download_end_exclusive"] == "2026-01-01"


def test_paper_snapshot_has_exactly_50_unique_securities():
    t = _load(PAPER_CFG)["tickers"]
    assert len(t) == 50
    assert len(set(t)) == 50


def test_paper_snapshot_includes_all_paper_example_tickers():
    t = _load(PAPER_CFG)["tickers"]
    for s in PAPER_EXAMPLES:
        assert s in t, f"{s} missing from the paper snapshot"
    assert set(_load(PAPER_CFG)["paper_example_tickers_required"]) == set(PAPER_EXAMPLES)


def test_paper_snapshot_reflects_the_2025_official_changes():
    t = _load(PAPER_CFG)["tickers"]
    for s in EXITS_2025:
        assert s not in t, f"{s} exited Nifty 50 in 2025 and must be absent"
    for s in ENTRANTS_2025:
        if s == "ZOMATO":
            # renamed ETERNAL per NSE circular CML67420, effective 2025-02-24
            assert "ETERNAL" in t
        else:
            assert s in t, f"{s} entered Nifty 50 in 2025 and must be present"


def test_paper_snapshot_provenance_lists_official_sources():
    prov = _load(PAPER_CFG)["provenance"]
    assert "nseindia.com" in prov["primary_list_source"]
    dates = {e["date"] for e in prov["bracketing_events"]}
    assert {"2025-03-28", "2025-09-30", "2025-10-14", "2025-11-17"} <= dates
    for e in prov["bracketing_events"]:
        assert "nseindia.com" in e["source"] or "niftyindices.com" in e["source"]


def test_dummytatam_is_excluded_as_an_index_entity():
    cfg = _load(PAPER_CFG)
    assert "DUMMYTATAM" not in cfg["tickers"]
    excluded = cfg["provenance"]["synthetic_index_entities_excluded"]
    assert any(e["symbol"] == "DUMMYTATAM" for e in excluded)
    note = next(e for e in excluded if e["symbol"] == "DUMMYTATAM")
    assert "no price history" in note["reason"]


def test_tata_motors_line_uses_tmpv_with_isin_continuity():
    """Tata Motors PV keeps the original ISIN, so it is the same listed entity."""
    t = _load(PAPER_CFG)["tickers"]
    assert "TMPV" in t
    assert "DUMMYTATAM" not in t
    assert _load(PAPER_CFG)["provenance"]["primary_list_source"]


def test_paper_snapshot_does_not_overwrite_the_other_universes():
    for path in (MODERN_CFG, LEGACY_CFG):
        text = (REPO_ROOT / path).read_text()
        assert "PAPER_FIXED_NIFTY50_2025_11_04" not in text, path
    assert _load(MODERN_CFG)["tickers"] != _load(PAPER_CFG)["tickers"]
    assert _load(LEGACY_CFG)["tickers"] != _load(PAPER_CFG)["tickers"]


def test_paper_snapshot_allows_late_listings_without_backfill():
    cfg = _load(PAPER_CFG)
    assert cfg["allow_late_listings"] is True
    text = (REPO_ROOT / PAPER_CFG).read_text()
    assert "backfilled" in text and "synthesised" in text


# ---------------------------------------------- three-universe comparison

def test_universe_comparison_covers_all_paper_examples():
    import csv
    path = REPO_ROOT / "results" / "reproduction_recovery" / "universe_comparison.csv"
    if not path.is_file():
        pytest.skip("universe_comparison.csv not generated yet")
    rows = {r["security"]: r for r in csv.DictReader(path.open())}
    for s in PAPER_EXAMPLES:
        assert s in rows, s
        assert rows[s]["paper_snapshot_2025_11_04"] == "Y", s
    counts = {
        "modern": sum(1 for r in rows.values() if r["modern_reconstructed"] == "Y"),
        "legacy": sum(1 for r in rows.values() if r["legacy_user_supplied"] == "Y"),
        "paper": sum(1 for r in rows.values() if r["paper_snapshot_2025_11_04"] == "Y"),
    }
    assert counts == {"modern": 50, "legacy": 50, "paper": 50}


def test_tcs_absent_from_legacy_but_present_in_paper_snapshot():
    """The contradiction that forced this whole exercise."""
    legacy = _load(LEGACY_CFG)["tickers"]
    paper = _load(PAPER_CFG)["tickers"]
    assert "TCS" not in legacy
    assert "TCS" in paper


# ------------------------------------------ paper-snapshot perf configs

@pytest.mark.parametrize("name", [
    "paper_snapshot_2025_unadjusted_perf",
    "paper_snapshot_2025_adjusted_perf",
])
def test_paper_perf_config(name):
    cfg = _load(f"configs/reproduction_search/{name}.yaml")
    data = cfg["data"]
    assert data["universe_id"] == "PAPER_FIXED_NIFTY50_2025_11_04_CANDIDATE"
    assert data["source_level"] == "daily"
    assert data["resample_to_daily"] is False
    assert data["sequence_length"] == 30
    assert data["expected_available_tickers"] == 50
    assert data["one_model_per_stock"] is True
    assert data["allow_late_listings"] is True
    assert cfg["evaluation"]["walk_forward"]["folds"] == "paper_exact"
    for model in ("attention_lstm", "lstm"):
        m = cfg["models"][model]
        assert m["max_epochs"] == 100, name
        assert m["patience"] == 10, name
        assert m["restore_best_checkpoint"] is True, name
        assert "epochs" not in m, f"{name}:{model} must not pin the 3-epoch cap"


def test_paper_perf_variants_differ_only_in_adjustment():
    u = _load("configs/reproduction_search/paper_snapshot_2025_unadjusted_perf.yaml")
    a = _load("configs/reproduction_search/paper_snapshot_2025_adjusted_perf.yaml")
    assert u["data"]["auto_adjust"] is False
    assert a["data"]["auto_adjust"] is True
    for key in ("sequence_length", "expected_available_tickers", "one_model_per_stock"):
        assert u["data"][key] == a["data"][key]
    assert u["models"] == a["models"]


# ------------------------------------------------- real dataset (opt-in)

@requires_paper_dataset
def test_paper_dataset_downloaded_all_fifty():
    files = list((PAPER_ROOT / "unadjusted" / "csv").glob("*.csv"))
    assert len(files) == 50, f"expected 50 CSVs, found {len(files)}"
    adj = list((PAPER_ROOT / "adjusted" / "csv").glob("*.csv"))
    assert len(adj) == 50


@requires_paper_dataset
def test_paper_dataset_canonical_columns():
    import pandas as pd
    for variant in ("unadjusted", "adjusted"):
        f = PAPER_ROOT / variant / "csv" / "TCS.csv"
        cols = list(pd.read_csv(f, nrows=0).columns)
        assert cols == ["Date", "Open", "High", "Low", "Close", "Volume"], (variant, cols)


@requires_paper_dataset
def test_paper_examples_cover_2022_and_2023():
    import pandas as pd
    for t in PAPER_EXAMPLES:
        d = pd.read_parquet(PAPER_ROOT / "unadjusted" / "parquet" / f"{t}.parquet")
        d["Date"] = pd.to_datetime(d["Date"])
        years = d["Date"].dt.year
        assert int((years == 2022).sum()) > 200, t
        assert int((years == 2023).sum()) > 200, t


@requires_paper_dataset
def test_late_listings_are_not_backfilled():
    """JIOFIN has no pre-2023 history and must not be invented."""
    import pandas as pd
    d = pd.read_parquet(PAPER_ROOT / "unadjusted" / "parquet" / "JIOFIN.parquet")
    d["Date"] = pd.to_datetime(d["Date"])
    assert d["Date"].min() >= pd.Timestamp("2023-01-01")
    assert d["Date"].is_monotonic_increasing
    assert not d["Date"].duplicated().any()


@requires_paper_dataset
def test_paper_workbooks_have_fifty_sheets():
    from openpyxl import load_workbook
    for variant in ("unadjusted", "adjusted"):
        f = PAPER_ROOT / "workbooks" / variant / "Nifty-50.xlsx"
        wb = load_workbook(f, read_only=True)
        assert len(wb.sheetnames) == 50, (variant, len(wb.sheetnames))
        for s in PAPER_EXAMPLES:
            assert s in wb.sheetnames, (variant, s)
        wb.close()


@requires_paper_dataset
def test_paper_dataset_root_is_separate_from_the_other_two():
    data = Path(_paper_root()).parent
    legacy = data / "yfinance_legacy_nifty50_2000_2025"
    modern = data / "yfinance_daily_2000_2025"
    assert PAPER_ROOT.name == "yfinance_paper_snapshot_2025_11_04"
    for other in (legacy, modern):
        if other.is_dir():
            assert PAPER_ROOT.resolve() != other.resolve()
